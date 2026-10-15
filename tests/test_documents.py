"""Word/PDF import: text order, image markers, filtering (CPU only, no models)."""

import io

import pytest
from PIL import Image

from mediaforge.text import documents as D


def _png(w: int, h: int, color=(200, 40, 40)) -> bytes:
    b = io.BytesIO()
    Image.new("RGB", (w, h), color).save(b, "PNG")
    return b.getvalue()


@pytest.fixture
def docx_file(tmp_path):
    import docx

    d = docx.Document()
    d.add_heading("Bảo mật mật khẩu", level=1)
    d.add_paragraph("Mật khẩu mạnh cần dài.")
    d.add_picture(io.BytesIO(_png(400, 300)))
    d.add_paragraph("Hãy dùng trình quản lý mật khẩu.")
    d.add_picture(io.BytesIO(_png(40, 40, (0, 0, 255))))  # icon: dropped
    d.add_picture(io.BytesIO(_png(400, 300)))  # duplicate of image 1: dropped
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "MFA"
    t.rows[0].cells[1].text = "Bật cho mọi tài khoản"
    d.add_picture(io.BytesIO(_png(500, 500, (20, 160, 20))))
    d.add_paragraph("Kết thúc.")
    p = tmp_path / "in.docx"
    d.save(p)
    return p


def test_docx_text_images_and_markers_in_order(docx_file, tmp_path):
    r = D.parse(docx_file, tmp_path / "out")
    assert r["kind"] == "word" and r["title"] == "Bảo mật mật khẩu"
    t = r["text"]
    assert len(r["images"]) == 2
    order = ["Mật khẩu mạnh", "[ảnh 1]", "trình quản lý", "MFA — Bật cho mọi tài khoản.", "[ảnh 2]", "Kết thúc."]
    pos = [t.index(s) for s in order]
    assert pos == sorted(pos)
    assert all(Image.open(im["path"]).format == "PNG" for im in r["images"])
    assert any("nhỏ" in w for w in r["warnings"])


def test_pdf_text_then_page_images(tmp_path):
    from pypdf import PdfReader, PdfWriter

    # A text page (hand-written minimal PDF) followed by an image page that Pillow writes.
    text_pdf = tmp_path / "t.pdf"
    stream = b"BT /F1 18 Tf 72 720 Td (Hello security team. This is page one.) Tj ET"
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
            b"<< /Length %d >>stream\n" % len(stream) + stream + b"\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offs = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offs.append(len(out))
        out += b"%d 0 obj\n" % i + o + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1) + b"".join(b"%010d 00000 n \n" % o for o in offs)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    text_pdf.write_bytes(out)
    img_pdf = tmp_path / "i.pdf"
    Image.new("RGB", (600, 400), (10, 120, 200)).save(img_pdf, "PDF")
    w = PdfWriter()
    for f in (text_pdf, img_pdf):
        for pg in PdfReader(f).pages:
            w.add_page(pg)
    doc = tmp_path / "doc.pdf"
    w.write(doc)

    r = D.parse(doc, tmp_path / "out")
    assert r["kind"] == "pdf" and len(r["images"]) == 1
    assert r["text"].index("Hello security team") < r["text"].index("[ảnh 1]")


def test_scanned_pdf_is_rejected(tmp_path):
    p = tmp_path / "scan.pdf"
    Image.new("RGB", (600, 800), "white").save(p, "PDF")
    with pytest.raises(D.DocumentError, match="scan"):
        D.parse(p, tmp_path / "out")


def test_unsupported_format(tmp_path):
    p = tmp_path / "x.pptx"
    p.write_bytes(b"x")
    with pytest.raises(D.DocumentError):
        D.parse(p, tmp_path / "out")


def test_unwrap_joins_pdf_lines():
    raw = "Mật khẩu mạnh cần\ndài và khó đoán.\nCâu thứ hai.\n\n- gạch đầu dòng\n- thứ hai\nsecu-\nrity"
    assert D._unwrap(raw) == ("Mật khẩu mạnh cần dài và khó đoán.\nCâu thứ hai.\n\n- gạch đầu dòng\n\n"
                              "- thứ hai security")


def test_markers_split_and_strip():
    text = "Mở đầu.\n\n[ảnh 1]\n\nThân bài.\n\n[Image 2]\n\nKết."
    assert D.split_markers(text) == [(None, "Mở đầu."), (1, "Thân bài."), (2, "Kết.")]
    assert D.strip_markers(text) == "Mở đầu.\n\nThân bài.\n\nKết."
    assert D.split_markers("[ảnh 3] Chỉ có ảnh.") == [(3, "Chỉ có ảnh.")]
    assert not D.has_markers("không có ảnh")


def test_unwrap_separates_pdf_headings():
    raw = ("Tiêu đề ngắn\nMột đoạn văn rất dài kéo dài qua nhiều dòng của trang PDF và vẫn chưa hết câu\n"
           "tiếp tục ở đây.")
    assert D._unwrap(raw) == ("Tiêu đề ngắn.\n\nMột đoạn văn rất dài kéo dài qua nhiều dòng của trang PDF và vẫn "
                              "chưa hết câu tiếp tục ở đây.")


# ---------------------------------------------------------------- shredded PDFs (docs/SDD-pdf-import-fix.md)

def _pdf_bytes(stream: bytes) -> bytes:
    """A one-page PDF with Helvetica as /F1 and the given content stream."""
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
            b"<< /Length %d >>stream\n" % len(stream) + stream + b"\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offs = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offs.append(len(out))
        out += b"%d 0 obj\n" % i + o + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1) + b"".join(b"%010d 00000 n \n" % o for o in offs)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    return out


def _shredded_pdf(path, lines: list[tuple[float, float, str]]) -> None:
    """Like Google Docs/Skia: a flipped page, then every word and every space in its own BT…ET with
    `1 0 0 -1 0 base Tm  x -d Td (w) Tj`. lines = [(y from the bottom, font size in points, text)]."""
    ops = [b"1 0 0 -1 0 792 cm"]
    for y, size, text in lines:
        fs = size / 0.75
        ops.append(b"q .75 0 0 .75 72 %.2f cm" % (792 - y - size))
        x = 0.0
        for i, word in enumerate(text.split(" ")):
            for w in ([" "] if i else []) + [word]:
                ops.append(b"BT /F1 %.2f Tf 1 0 0 -1 0 %.2f Tm %.2f %.2f Td (%s) Tj ET"
                           % (fs, fs * 1.2, x, -fs * 0.9, w.encode("latin-1")))
                x += fs * (0.3 if w == " " else 0.55 * len(w))
        ops.append(b"Q")
    path.write_bytes(_pdf_bytes(b"\n".join(ops)))


LONG1 = "Strong passwords need to be long and hard to guess for every account"
LONG2 = "and they should never be shared with anyone at all."
LONG3 = "A second paragraph starts here after the normal spacing between paragraphs"
LONG4 = "and continues on its second line until it ends."


def test_pdf_shredded_words_become_paragraphs(tmp_path):
    p = tmp_path / "gdocs.pdf"
    _shredded_pdf(p, [(740, 15, "Security Basics"),
                      (712, 12, LONG1), (694.8, 12, LONG2),
                      (671.6, 12, LONG3), (654.4, 12, LONG4)])
    assert D._shredded(__import__("pypdf").PdfReader(str(p)).pages[0].extract_text())
    blocks = D.parse(p, tmp_path / "out")["text"].split("\n\n")
    assert blocks == ["Security Basics.", f"{LONG1} {LONG2}", f"{LONG3} {LONG4}"]


def test_pdf_shredded_short_page(tmp_path):
    p = tmp_path / "cover.pdf"
    _shredded_pdf(p, [(700, 20, "Annual Security Report"), (660, 12, "Prepared by the security team")])
    blocks = D.parse(p, tmp_path / "out")["text"].split("\n\n")
    assert all(len(b.split()) > 1 for b in blocks)
    assert blocks[0] == "Annual Security Report."


def test_shredded_threshold():
    assert not D._shredded("A normal line of prose with words\nand another normal line here\nand a third one")
    assert D._shredded("Phần\n \n4.\n \nTrách\n \nnhiệm")
    assert not D._shredded("Một\nHai")  # too few lines to judge


def test_geometric_failure_falls_back(tmp_path, monkeypatch):
    p = tmp_path / "gdocs.pdf"
    _shredded_pdf(p, [(712, 12, LONG1), (694.8, 12, LONG2)])

    def boom(page):
        raise RuntimeError("visitor broke")
    monkeypatch.setattr(D, "_geometric", boom)
    text = D.parse(p, tmp_path / "out")["text"]
    assert "Strong passwords need" in text and "\n\nneed\n\n" not in text


def test_unwrap_numbered_headings():
    raw = ("Phần 4. Trách nhiệm xã hội\n4.1. Hồ sơ CSR\n"
           "PNJ là doanh nghiệp vàng duy nhất trong nhóm bốn công ty có hệ thống công bố phát triển bền vững\n"
           "liên tục từ 2019.")
    assert D._unwrap(raw).split("\n\n") == [
        "Phần 4. Trách nhiệm xã hội.", "4.1. Hồ sơ CSR.",
        ("PNJ là doanh nghiệp vàng duy nhất trong nhóm bốn công ty có hệ thống công bố phát triển bền vững "
         "liên tục từ 2019.")]
    assert not D.HEADING_RE.match("2025 là năm bản lề")
